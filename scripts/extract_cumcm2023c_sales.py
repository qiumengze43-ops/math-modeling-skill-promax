"""Extract category-day sales from the public 2023 CUMCM C attachments."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Iterable


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _column(headers: list[str], name: str) -> int:
    try:
        return headers.index(name)
    except ValueError as exc:
        raise ValueError(f"Missing required column {name!r}; found {headers!r}") from exc


def _date_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _text(value).split(" ", 1)[0]


def read_category_map(attachment1: Path) -> dict[str, str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading XLSX attachments requires openpyxl in the active environment") from exc

    workbook = load_workbook(attachment1, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(rows)]
    item_column = _column(headers, "单品编码")
    category_column = _column(headers, "分类名称")
    mapping: dict[str, str] = {}
    for row in rows:
        item = _text(row[item_column])
        category = _text(row[category_column])
        if item and category:
            mapping[item] = category
    workbook.close()
    if not mapping:
        raise ValueError("Attachment 1 did not yield any item-to-category mappings")
    return mapping


def extract_sales(attachment1: Path, attachment2: Path, output_csv: Path) -> dict[str, object]:
    category_by_item = read_category_map(attachment1)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading XLSX attachments requires openpyxl in the active environment") from exc

    workbook = load_workbook(attachment2, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(rows)]
    date_column = _column(headers, "销售日期")
    item_column = _column(headers, "单品编码")
    quantity_column = _column(headers, "销量(千克)")
    daily: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    row_count = 0
    matched_rows = 0
    for row in rows:
        row_count += 1
        item = _text(row[item_column])
        category = category_by_item.get(item)
        if not category:
            continue
        try:
            quantity = float(row[quantity_column])
        except (TypeError, ValueError):
            continue
        day = _date_text(row[date_column])
        if not day:
            continue
        daily[day][category] += quantity
        matched_rows += 1
    workbook.close()

    categories = sorted({category for values in daily.values() for category in values})
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["销售日期", *categories])
        for day in sorted(daily):
            writer.writerow([day, *(round(daily[day].get(category, 0.0), 6) for category in categories)])
    return {"sales_rows": row_count, "matched_rows": matched_rows, "days": len(daily), "categories": categories}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attachment1", type=Path, required=True)
    parser.add_argument("--attachment2", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(extract_sales(args.attachment1, args.attachment2, args.output))


if __name__ == "__main__":
    main()
